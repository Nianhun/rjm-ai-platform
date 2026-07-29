from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("瑞吉明生物医药AI智能系统详细设计文档.docx")


BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(11, 37, 69)
MUTED = RGBColor(89, 89, 89)
LIGHT_GRAY = "F2F4F7"
LIGHT_BLUE = "E8EEF5"
CALLOUT = "F4F6F9"
BORDER = "C8CED8"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for w in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(w))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Pt(widths[idx] / 20)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_borders(table, color=BORDER):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_run_font(run, size=None, color=None, bold=None, italic=None, font="Microsoft YaHei"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_para(doc, text="", style=None, bold=False, color=None, size=None, align=None, after=6, before=0):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size, color=color, bold=bold)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.167
        r = p.add_run(item)
        set_run_font(r, size=10.5)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.167
        r = p.add_run(item)
        set_run_font(r, size=10.5)


def add_callout(doc, title, body, fill=CALLOUT):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    set_borders(table, "D8DEE8")
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    rt = p.add_run(title)
    set_run_font(rt, size=10.5, color=INK, bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    rb = p2.add_run(body)
    set_run_font(rb, size=10.5)
    add_para(doc, "", after=3)


def add_table(doc, headers, rows, widths, header_fill=LIGHT_GRAY):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    set_borders(table)
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_shading(hdr[i], header_fill)
        p = hdr[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        set_run_font(r, size=9.5, bold=True, color=INK)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            r = p.add_run(str(value))
            set_run_font(r, size=9.2)
    set_table_geometry(table, widths)
    add_para(doc, "", after=4)
    return table


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        s = styles[name]
        s.font.name = "Microsoft YaHei"
        s._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        s.font.size = Pt(size)
        s.font.color.rgb = color
        s.font.bold = True
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.line_spacing = 1.1

    for list_name in ["List Bullet", "List Number"]:
        s = styles[list_name]
        s.font.name = "Microsoft YaHei"
        s._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        s.font.size = Pt(10.5)
        s.paragraph_format.space_after = Pt(4)
        s.paragraph_format.line_spacing = 1.167


def add_footer(section):
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("烟台瑞吉明生物医药AI智能系统详细设计文档")
    set_run_font(run, size=9, color=MUTED)


def add_title_page(doc):
    add_para(doc, "详细设计文档", bold=True, size=23, color=INK, after=4)
    add_para(doc, "烟台瑞吉明生物医药AI智能系统", size=15, color=MUTED, after=18)

    rows = [
        ("建设单位", "烟台瑞吉明公司"),
        ("文档定位", "第一版系统详细设计与实施计划"),
        ("设计范围", "知识库、文献分析、原料图谱、AI配方推荐、实验反馈学习、中试模拟原型；真实设备控制暂列为接口与数据回流设计"),
        ("技术路线", "前期 Python 快速验证；知识底座复用 Yuxi；后续管理系统采用 Java"),
        ("形成日期", str(date.today())),
    ]
    add_table(doc, ["项目", "内容"], rows, [1800, 7560], LIGHT_GRAY)

    add_callout(
        doc,
        "设计原则",
        "系统第一阶段不追求一次性完成智能工厂闭环，而是先跑通可追溯的配方研发闭环：文献、专利、原料属性和实验记录进入知识库与图谱，AI基于原料关系生成配方推荐清单，配方工程师筛选后开展实验，实验反馈回流系统，成功配方再进入供应商采购和中试流程。",
        LIGHT_BLUE,
    )
    doc.add_page_break()


def add_toc(doc):
    doc.add_heading("目录", level=1)
    sections = [
        "1. 项目背景与建设目标",
        "2. 需求理解与范围边界",
        "3. 总体架构设计",
        "4. 技术路线与系统分工",
        "5. 核心业务闭环设计",
        "6. 文献数据分析平台详细设计",
        "7. AI研发建议模块详细设计",
        "8. 中试智能模拟测试平台详细设计",
        "9. 设备供应商与采购体系详细设计",
        "10. 虚实转换模块详细设计",
        "11. 化妆品已备案数据库设计",
        "12. 知识库、数据模型与知识图谱设计",
        "13. Agent工作流设计",
        "14. 接口与集成设计",
        "15. 安全、权限、审计与数据治理",
        "16. 部署方案与环境规划",
        "17. 项目里程碑与实施计划",
        "18. 风险、难点与应对策略",
        "19. 验收标准",
        "20. 后续演进路线",
    ]
    for item in sections:
        add_para(doc, item, after=2)
    doc.add_page_break()


def build():
    doc = Document()
    style_document(doc)
    add_footer(doc.sections[0])
    add_title_page(doc)
    add_toc(doc)

    doc.add_heading("1. 项目背景与建设目标", level=1)
    add_para(doc, "烟台瑞吉明公司提出的生物医药AI智能系统，核心诉求不是建设单点问答工具，而是围绕文献、专利、实验、中试、设备和工厂数据形成可追溯、可验证、可迭代的研发智能系统。")
    add_para(doc, "参考材料中有两条关键启发：一是 OpenAI AI 化学家展示了“提出假设、设计实验、执行验证、根据结果迭代”的科研闭环；二是材料大模型综述指出，真正困难的是把文献抽取、结构化表示、性能预测、工具调用和实验闭环接稳。瑞吉明项目应以这两点作为第一性原则。")
    add_callout(doc, "建设目标", "第一版先跑通“文献/专利/原料知识入库 - 原料属性与关系图谱 - AI配方推荐List - 工程师经验筛选 - 真实实验反馈 - 配方效果学习 - 供应商采购衔接”的闭环；后续再扩展 Java 管理系统、真实设备接口和智能工厂执行。")
    add_para(doc, "典型目标场景是：配方工程师提出“我要一个保湿方向的配方”，AI 能够基于图谱中的原材料属性、功效、相容性、禁忌关系、用量范围、法规限制、供应商信息和历史实验反馈，输出多套配方推荐 List。工程师依据经验过滤其中一部分，剩余配方进入真实实验；实验结果反馈给系统后，AI 更新原料关系权重、配方评分和推荐策略。")

    doc.add_heading("2. 需求理解与范围边界", level=1)
    add_table(
        doc,
        ["需求板块", "原始诉求", "第一版设计定位"],
        [
            ("文献数据分析平台", "搜集全球新材料、生物医药相关文献，定期自动检索归纳，专利推送与分析。", "作为MVP主入口，优先建设文献/专利采集、知识库入库、结构化抽取、综述生成和证据追溯。"),
            ("临床数据文献", "找到标准化格式及字段，精确抓取核心数据，英文文献需要翻译。", "优先做字段模板、英文翻译、实体关系抽取和数据质检。"),
            ("科研文献", "内容冗杂，字段不固定，看情况归纳整理。", "采用RAG摘要、主题聚类、研究方向归纳和人工确认机制。"),
            ("中试实验平台", "模拟实验室到小试、中试、大生产的放大过程，发现问题并优化方案。", "先做模拟原型和反馈闭环，不直接控制真实设备。"),
            ("设备供应商体系", "供应商端和采购端上传设备数据，中试时推荐设备。", "先设计设备主数据、选型规则和推荐流程。"),
            ("原材料供应商体系", "成功配方需要找到原材料供应商并进行采购。", "将原料主数据、供应商、价格、交期、资质、批次质量纳入配方闭环。"),
            ("虚实转换模块", "对接工厂设备接口，收集数据，输出方案给智能工厂执行。", "第一版只做接口规范、数据采集和方案回写设计。"),
            ("化妆品备案数据库", "调用国家局数据接口并分析。", "作为独立数据服务，纳入检索、比对和合规分析。"),
        ],
        [1600, 3800, 3960],
        LIGHT_BLUE,
    )
    add_para(doc, "第一版明确不直接控制真实工厂设备，不承诺自动完成生产放大决策，不把AI输出作为无需人工审批的最终结论。所有配方推荐、研发建议和中试方案均需保留证据链、置信度、人工确认状态和线下验证结果。")

    doc.add_heading("3. 总体架构设计", level=1)
    add_para(doc, "系统采用“Python AI能力验证层 + Yuxi知识智能底座 + Java业务管理系统 + 工厂/外部数据接口层”的分层架构。")
    add_table(
        doc,
        ["层级", "核心职责", "主要技术/系统"],
        [
            ("数据接入层", "接入文献、专利、临床数据、科研文献、原料资料、供应商资料、设备参数、实验记录、备案库数据。", "Python采集脚本、API连接器、文件上传、定时任务"),
            ("知识智能层", "文档解析、分块、向量检索、知识图谱、引用追溯、Agent编排。", "Yuxi、Milvus、Neo4j、PostgreSQL、MinIO、LangGraph"),
            ("AI算法层", "结构化抽取、翻译、归纳、配方推荐、假设生成、中试模拟、结果分析与迭代优化。", "Python、LLM、领域规则、统计模型、推荐算法、可插拔仿真模块"),
            ("业务管理层", "项目、用户、权限、任务、原料、供应商、采购、设备、实验反馈、审批和审计。", "Java Spring Boot、关系数据库、流程引擎"),
            ("集成接口层", "对接专利/备案/工厂设备/智能工厂/企业系统。", "REST API、消息队列、文件交换、工业协议适配器"),
        ],
        [1500, 4400, 3460],
        LIGHT_GRAY,
    )
    add_para(doc, "架构上，Yuxi 不替代 Java 管理系统，而是作为企业科研知识和智能体底座。Java 系统负责业务对象和流程闭环，Python 负责AI算法快速实验，二者通过标准API共享知识检索、原料图谱、配方推荐结果、实验反馈和中试模拟结果。")

    doc.add_heading("4. 技术路线与系统分工", level=1)
    add_table(
        doc,
        ["技术单元", "近期职责", "长期职责"],
        [
            ("Python", "快速验证文献抽取、翻译、专利解析、原料关系建模、配方推荐、实验反馈学习、中试模拟算法。", "沉淀为AI服务、批处理任务、推荐模型、模型评估与算法迭代平台。"),
            ("Yuxi", "承载知识库搭建、文档解析、RAG检索、图谱抽取、Agent工作台。", "作为科研知识中台，为Java系统和AI服务提供检索、引用和智能体能力。"),
            ("Java管理系统", "第二阶段开始建设项目、任务、权限、流程、供应商、采购、设备管理。", "成为企业级生产系统入口，负责审计、审批、数据治理和业务协同。"),
            ("数据库与中间件", "PostgreSQL保存业务与抽取结果，Milvus做向量检索，Neo4j做知识图谱，MinIO存文件。", "按企业数据治理要求扩展备份、权限、审计、生命周期管理。"),
            ("模型服务", "使用兼容OpenAI接口的大模型做抽取、总结、推理和工具调用。", "引入领域微调、评估集、模型路由、私有化部署和成本控制。"),
        ],
        [1800, 3900, 3660],
        LIGHT_BLUE,
    )

    doc.add_heading("5. 核心业务闭环设计", level=1)
    add_para(doc, "系统核心闭环如下：")
    add_numbered(doc, [
        "知识进入：文献、专利、备案数据、原料资料、供应商资料、实验记录、设备参数进入知识库和业务数据库。",
        "知识转化：AI抽取原料属性、功效、适用场景、相容性、禁忌关系、用量范围、法规限制、实验条件、工艺参数和结果指标。",
        "配方推荐：配方Agent基于原料图谱、历史实验、目标功效和约束条件生成配方推荐 List，并给出推荐理由、风险点和证据链。",
        "人工筛选：配方工程师依据经验、成本、供应链、法规和工艺可行性过滤候选配方，选择进入实验的方案。",
        "真实验证：实验人员按筛选后的方案开展实验，回填肤感、稳定性、功效评价、成本、异常和结论。",
        "反馈学习：系统比较推荐评分与实验结果，更新原料关系权重、配方相似案例、成功/失败模式和下一轮推荐策略。",
        "采购衔接：验证可行的配方进入原料供应商匹配、询价、样品采购和批次质量跟踪流程。",
    ])
    add_callout(doc, "闭环评价标准", "AI不以“回答得像”为成功标准，而以“配方有证据、原料关系可解释、工程师可筛选、实验可验证、反馈能学习、采购能衔接”为成功标准。")

    doc.add_heading("5.1 保湿配方推荐示例流程", level=2)
    add_table(
        doc,
        ["步骤", "系统动作", "人工动作/结果"],
        [
            ("提出目标", "用户输入：需要一个保湿方向配方，可补充肤感、成本、剂型、法规、原料黑名单等约束。", "配方工程师确认目标和约束。"),
            ("图谱推理", "系统检索保湿相关原料，分析透明质酸钠、甘油、泛醇、神经酰胺、甜菜碱等原料的属性、协同关系、冲突关系和历史反馈。", "工程师查看证据链和推荐理由。"),
            ("生成List", "输出多套候选配方，每套包含原料、建议比例范围、核心功效逻辑、工艺注意点、风险和供应商可得性。", "工程师依据经验筛掉不合适方案。"),
            ("真实实验", "系统把保留方案转为实验记录模板，跟踪批次、参数和评价指标。", "实验人员开展实验并回填结果。"),
            ("反馈进化", "系统将成功/失败、稳定性、肤感、成本和异常原因写回知识图谱与推荐评分。", "后续同类需求优先推荐更可靠组合。"),
            ("采购衔接", "可行配方进入原料供应商匹配、询价和样品采购。", "采购部门执行采购并回填供应质量。"),
        ],
        [1200, 5000, 3160],
        LIGHT_BLUE,
    )

    doc.add_heading("6. 文献数据分析平台详细设计", level=1)
    doc.add_heading("6.1 功能组成", level=2)
    add_bullets(doc, [
        "文献源管理：维护 PubMed、Google Scholar、Crossref、专利库、企业自有资料等数据源配置。",
        "定时检索任务：按关键词、疾病领域、材料/药物方向、作者、机构、专利权人等条件周期检索。",
        "文献入库：将PDF、DOCX、网页、摘要、表格和图片解析为Markdown和结构化元数据。",
        "英文翻译：对标题、摘要、实验段落、表格字段做可追溯翻译，保留原文和译文。",
        "结构化抽取：抽取研究对象、实验条件、样本量、指标、结果、结论、限制、专利权利要求等字段。",
        "文献综述生成：按主题生成综述，自动插入站内链接、下载链接和引用来源。",
        "专利分析：拆解技术方案、权利要求、保护范围、相似专利和潜在规避方向。",
    ])
    doc.add_heading("6.2 临床文献字段模板", level=2)
    add_table(
        doc,
        ["字段类别", "建议字段", "说明"],
        [
            ("基础信息", "标题、作者、期刊、年份、DOI、国家、机构", "用于来源追溯和检索筛选。"),
            ("研究设计", "研究类型、入排标准、样本量、分组、随机/盲法", "优先抽取标准化字段。"),
            ("干预/治疗", "药物/材料/方法、剂量、频次、周期、联合方案", "支持英文原文和中文译文对照。"),
            ("终点指标", "主要终点、次要终点、安全性指标、统计方法", "必须绑定具体数值、单位和统计显著性。"),
            ("结果结论", "疗效、安全性、不良事件、作者结论、局限性", "保留证据片段和置信度。"),
        ],
        [1500, 3600, 4260],
        LIGHT_GRAY,
    )
    doc.add_heading("6.3 科研文献处理策略", level=2)
    add_para(doc, "科研文献字段不稳定，应采用“固定元数据 + 动态实体关系 + 主题归纳”的模式。系统先抽取材料/药物对象、合成或制备过程、处理参数、性能指标、测试条件和结论，再由人工确认是否进入长期结构化数据库。")

    doc.add_heading("7. AI研发建议模块详细设计", level=1)
    add_para(doc, "AI研发建议模块面向研发人员，目标是把文献知识、专利信息、实验记录和企业目标转化为可讨论、可验证、可证伪的研发建议。")
    add_table(
        doc,
        ["子模块", "输入", "输出"],
        [
            ("研发方向识别", "文献主题、专利趋势、企业关注方向", "候选研发方向、热点变化、竞争态势"),
            ("假设生成", "知识库证据、已有实验数据、约束条件", "可证伪假设、预期结果、关键验证指标"),
            ("实验建议", "假设、实验资源、设备可用性", "小试方案、参数范围、对照组、风险点"),
            ("证据链管理", "RAG检索结果、图谱关系、原文片段", "引用来源、证据等级、人工确认状态"),
            ("评估与迭代", "实验反馈、预测值、实测值", "误差分析、下一轮建议、模型评估记录"),
        ],
        [1700, 3600, 4060],
        LIGHT_BLUE,
    )
    add_callout(doc, "研发建议约束", "系统输出必须同时包含建议、依据、假设、验证方法和失败条件。没有证据来源的建议不得进入中试方案。")

    doc.add_heading("7.1 AI配方推荐模块详细设计", level=2)
    add_para(doc, "配方推荐模块是项目最终效果的核心能力。它面向配方工程师，不直接替代工程师决策，而是根据原料图谱、原料属性、历史实验反馈和供应链信息，生成可解释、可筛选、可实验的候选配方 List。")
    add_table(
        doc,
        ["输入", "处理逻辑", "输出"],
        [
            ("功效目标", "识别保湿、舒缓、修护、美白、抗皱等目标，并映射到功效原料和评价指标。", "目标功效拆解与评价指标清单"),
            ("配方约束", "解析剂型、成本、法规、禁用原料、肤感、稳定性、供应商偏好等限制。", "硬约束和软约束"),
            ("原料图谱", "检索原料属性、功效、相容性、协同/拮抗关系、用量范围、风险和历史表现。", "候选原料池与关系证据"),
            ("历史反馈", "召回相似配方的成功/失败实验，计算原料组合的经验权重。", "推荐评分和风险评分"),
            ("供应链数据", "检查原料供应商、价格、交期、资质、样品可得性。", "供应可行性和采购建议"),
        ],
        [1700, 4300, 3360],
        LIGHT_GRAY,
    )
    add_table(
        doc,
        ["配方推荐List字段", "说明"],
        [
            ("配方编号", "例如 FORM-MOIST-20260727-001，便于实验和反馈追踪。"),
            ("目标功效", "如保湿、长效锁水、屏障修护等。"),
            ("原料组成", "原料名称、建议比例范围、角色：主功效、辅助保湿、增稠、防腐、肤感调节等。"),
            ("推荐理由", "基于图谱关系、文献证据、历史实验和专家规则生成。"),
            ("风险提示", "相容性、稳定性、法规、刺激性、成本、供应风险。"),
            ("证据链", "关联文献、专利、原料资料、实验记录、备案产品或工程师反馈。"),
            ("实验建议", "建议小试步骤、观察指标、对照组和失败判据。"),
            ("供应商建议", "可采购供应商、价格区间、交期、资质和样品状态。"),
            ("推荐评分", "综合功效匹配、实验成功率、风险、成本和供应链可行性。"),
        ],
        [2500, 6860],
        LIGHT_BLUE,
    )

    doc.add_heading("8. 中试智能模拟测试平台详细设计", level=1)
    add_para(doc, "中试平台第一版定位为“模拟与反馈系统”，不是直接控制生产设备的执行系统。它重点解决从实验室、小试到中试放大过程中的参数传递、风险识别、设备匹配、反馈优化。")
    add_table(
        doc,
        ["模块", "主要能力", "关键数据"],
        [
            ("小试数据管理", "录入配方、工艺、环境、产率、质量、异常记录。", "配方、批次、温度、压力、时间、pH、转速、收率、质量指标"),
            ("放大模拟", "根据小试参数和目标规模推演中试参数范围。", "放大倍数、传热传质、混合、能耗、设备能力、约束规则"),
            ("问题识别", "识别放大过程中可能出现的风险。", "过热、混合不均、沉淀、污染、稳定性、成本、设备瓶颈"),
            ("优化方案", "输出多套中试方案并排序。", "推荐参数、设备组合、风险等级、验证步骤、预计成本"),
            ("反馈学习", "线下实验结果回填后做预测-实测差异分析。", "实测数据、异常原因、成功/失败标签、修正建议"),
        ],
        [1600, 4300, 3460],
        LIGHT_GRAY,
    )
    add_para(doc, "算法原型可以先从规则、统计模型和轻量机器学习开始，不急于引入复杂数字孪生。随着真实反馈数据累积，再逐步引入更精细的机理模型、仿真模型和参数优化算法。")

    doc.add_heading("9. 设备供应商与采购体系详细设计", level=1)
    add_para(doc, "供应商与采购体系包括设备供应商和原材料供应商两类。对于配方研发闭环，原材料供应商体系尤为关键：AI推荐出的可行配方，需要进一步匹配可采购、资质合规、质量稳定的原料来源。")
    add_table(
        doc,
        ["角色", "能力", "权限边界"],
        [
            ("原料供应商端", "上传原料规格、INCI名称、COA、MSDS、报价、最小采购量、交期、样品状态。", "只能维护本供应商原料与资质数据。"),
            ("设备供应商端", "上传设备型号、参数、报价、交期、案例、认证、接口协议。", "只能维护本供应商数据，不能查看企业敏感中试数据。"),
            ("采购部门端", "维护采购需求、询价、比价、供应商评价、合同状态。", "查看原料/设备候选、采购流程和供应商资料。"),
            ("研发/中试端", "发起原料样品或设备需求，查看适配建议和测试结果。", "查看与项目相关的原料、设备信息和适配记录。"),
            ("管理员", "维护原料分类、设备分类、参数模板、供应商准入和权限。", "全局配置与审计。"),
        ],
        [1600, 4300, 3460],
        LIGHT_BLUE,
    )
    add_para(doc, "设备推荐采用“硬约束过滤 + 软指标排序”的方式。硬约束包括容量、材质、温压范围、洁净等级、接口能力；软指标包括价格、交期、历史评价、维护成本、适配案例和供应商响应速度。")
    add_para(doc, "原材料供应商推荐同样采用“硬约束过滤 + 软指标排序”。硬约束包括原料规格、法规资质、禁限用要求、最小采购量、交期和样品可得性；软指标包括价格、历史质量、批次稳定性、供应商服务、账期和与目标配方的适配经验。")

    doc.add_heading("10. 虚实转换模块详细设计", level=1)
    add_para(doc, "虚实转换模块是连接中试模拟和真实工厂的接口层。第一版只设计数据采集、指标映射、方案回写和人工审批，不直接下发控制指令。")
    add_bullets(doc, [
        "数据采集：从现有工厂设备接口采集运行状态、批次数据、报警、质量指标和能耗数据。",
        "指标映射：把设备字段映射到中试模型字段，例如温度、压力、流量、转速、时间、产量、合格率。",
        "方案回写：将中试方案转换为智能工厂可理解的工艺单、参数建议或任务草稿。",
        "人工审批：任何真实执行动作必须经过授权人员确认，系统保留审批记录。",
        "结果回流：实际执行数据回流至中试平台，用于验证模型和优化规则。",
    ])

    doc.add_heading("11. 化妆品已备案数据库设计", level=1)
    add_para(doc, "化妆品已备案数据库作为独立数据服务，负责调用国家局公开数据接口或导入备案数据，并对产品、成分、功效、企业、备案状态进行分析。")
    add_table(
        doc,
        ["能力", "说明"],
        [
            ("数据同步", "按接口限制定期同步备案产品、企业、成分、备案状态。"),
            ("成分分析", "分析成分组合、常见配伍、风险成分和功效宣称。"),
            ("竞品比对", "按品类、功效、成分、企业和时间维度做竞品分析。"),
            ("研发参考", "为新产品立项提供已备案产品证据和避坑建议。"),
            ("合规提示", "对功效宣称、禁限用成分和备案状态变化做提醒。"),
        ],
        [2200, 7160],
        LIGHT_GRAY,
    )

    doc.add_heading("12. 知识库、数据模型与知识图谱设计", level=1)
    add_para(doc, "第一阶段复用 Yuxi 已搭建的知识库能力：文件上传、解析、分块、PostgreSQL chunk 存储、Milvus 向量检索、Neo4j 图谱和 Agent 检索工具。")
    add_table(
        doc,
        ["对象", "关键字段"],
        [
            ("Document", "id、source_type、title、language、file_path、hash、status、created_at、kb_id"),
            ("ExtractedFact", "id、document_id、fact_type、subject、predicate、object、value、unit、evidence_span、confidence"),
            ("Ingredient", "id、name_cn、name_en、inci_name、category、functions、properties、usage_range、regulatory_limits、risk_tags"),
            ("IngredientRelation", "id、source_ingredient_id、target_ingredient_id、relation_type、strength、evidence_ids、feedback_weight"),
            ("FormulaCandidate", "id、goal、ingredients、ratio_ranges、recommendation_reason、risk_notes、score、status、evidence_ids"),
            ("FormulaExperiment", "id、formula_id、batch_no、actual_ratios、process_params、stability_result、efficacy_result、sensory_result、conclusion"),
            ("ResearchHypothesis", "id、topic、hypothesis、evidence_ids、expected_result、failure_condition、status"),
            ("ExperimentPlan", "id、project_id、hypothesis_id、scale、parameters、equipment_candidates、risk_level、approval_status"),
            ("ExperimentResult", "id、plan_id、batch_no、actual_parameters、metrics、deviations、conclusion、attachments"),
            ("Supplier", "id、supplier_type、name、qualification、contact、rating、delivery_cycle、quality_records"),
            ("RawMaterialSku", "id、ingredient_id、supplier_id、specification、price、moq、lead_time、coa_file、sample_status"),
            ("Equipment", "id、supplier_id、category、model、capacity、material、temperature_range、pressure_range、interfaces"),
        ],
        [2200, 7160],
        LIGHT_BLUE,
    )
    add_para(doc, "知识图谱建议优先抽取六类配方关系：原料-功效关系，原料-属性关系，原料-原料协同/冲突关系，原料-法规/风险关系，配方-实验结果关系，原料-供应商关系。文献-研究对象-指标关系、材料/药物-工艺-性能关系、专利-权利要求-技术点关系和实验方案-设备-结果关系作为扩展图谱继续保留。")
    add_callout(doc, "图谱进化机制", "实验反馈不是简单存档，而是会改变图谱中的关系权重。例如某个保湿组合多次实验稳定且肤感好，则增强相关协同关系；若出现分层、刺激性或稳定性问题，则增加风险标签并降低后续推荐权重。")

    doc.add_heading("13. Agent工作流设计", level=1)
    add_table(
        doc,
        ["Agent", "职责", "调用工具"],
        [
            ("文献检索Agent", "根据主题定期检索、去重、入库和生成摘要。", "检索API、Yuxi上传、知识库查询、翻译工具"),
            ("结构化抽取Agent", "从文献/专利中抽取字段、实体、关系和证据片段。", "RAG、OCR结果、JSON Schema校验、人工复核队列"),
            ("综述生成Agent", "基于多个文献证据生成综述和站内引用。", "query_kb、open_kb_document、引用管理"),
            ("配方推荐Agent", "根据目标功效生成配方推荐List，并解释原料选择、比例范围、风险和证据。", "原料图谱、历史实验、供应商库、法规规则、候选排序"),
            ("研发建议Agent", "提出可证伪研发假设和实验建议。", "知识图谱、历史实验、规则库、候选排序"),
            ("中试模拟Agent", "生成放大方案、设备建议、风险清单和反馈分析。", "实验数据库、设备库、模拟算法、审批接口"),
            ("采购推荐Agent", "为验证可行的配方匹配原材料供应商和采购建议。", "供应商库、RawMaterialSku、价格/交期/资质记录"),
            ("合规分析Agent", "分析备案库、专利、原料限制和合规风险。", "备案接口、专利库、法规规则库"),
        ],
        [1600, 4200, 3560],
        LIGHT_GRAY,
    )
    add_callout(doc, "人机协同机制", "每个Agent输出都应进入可审阅状态：草稿、待确认、已确认、已驳回、已用于方案。AI负责缩短分析时间，最终业务结论由授权人员确认。")

    doc.add_heading("14. 接口与集成设计", level=1)
    add_table(
        doc,
        ["接口", "方向", "说明"],
        [
            ("文献采集接口", "外部 -> Python/Yuxi", "按关键词、时间、领域拉取文献元数据和全文链接。"),
            ("知识库检索接口", "Java/Python -> Yuxi", "输入问题、过滤条件，返回片段、引用、文档链接和置信度。"),
            ("抽取结果接口", "Python -> Java", "回写结构化字段、证据、置信度和人工复核状态。"),
            ("研发建议接口", "Python/Yuxi -> Java", "返回假设、建议、证据链、风险和验证方案。"),
            ("配方推荐接口", "Python/Yuxi -> Java", "输入功效目标和约束，返回配方List、原料比例范围、证据链、风险和推荐评分。"),
            ("实验反馈接口", "Java -> Python/Yuxi", "回写配方工程师筛选结果、真实实验结果和失败原因，用于更新图谱权重和推荐策略。"),
            ("中试方案接口", "Java/Python -> Java", "创建、查询、审批、反馈中试方案和实验结果。"),
            ("设备数据接口", "供应商/工厂 -> Java", "上传设备参数、运行数据、报警和批次结果。"),
            ("备案数据库接口", "外部 -> Java/Python", "同步备案产品并提供查询、比对、分析能力。"),
        ],
        [1800, 1800, 5760],
        LIGHT_BLUE,
    )

    doc.add_heading("15. 安全、权限、审计与数据治理", level=1)
    add_bullets(doc, [
        "权限模型：超级管理员、管理员、研发人员、中试人员、采购人员、供应商用户、只读审计用户。",
        "知识库权限：按全局、部门、指定人员授权，敏感文献和实验数据默认最小授权。",
        "数据隔离：供应商端与企业研发数据隔离，供应商只能维护自己的设备信息。",
        "审计记录：记录文献入库、抽取修改、AI建议生成、方案审批、实验反馈、接口调用。",
        "AI输出治理：保存提示词版本、模型版本、检索证据、输出内容、人工确认状态。",
        "隐私与合规：临床文献和企业实验数据需脱敏处理，禁止把未授权数据发送到外部模型服务。",
    ])

    doc.add_heading("16. 部署方案与环境规划", level=1)
    add_table(
        doc,
        ["环境", "用途", "部署建议"],
        [
            ("开发环境", "功能开发和算法验证。", "Yuxi Docker Compose + Python本地/容器 + Java开发服务。"),
            ("测试环境", "集成测试、权限测试、抽取准确率评估。", "独立数据库和对象存储，使用脱敏样本。"),
            ("试运行环境", "真实小规模业务验证。", "接入部分真实文献、设备和实验数据，开启审计。"),
            ("生产环境", "企业正式使用。", "强密码、备份、监控、日志、模型密钥隔离、最小端口暴露。"),
        ],
        [1600, 3400, 4360],
        LIGHT_GRAY,
    )
    add_para(doc, "Yuxi 生产部署建议保留 PostgreSQL、Redis、MinIO、Milvus、Neo4j、API、Worker、Web 服务。Java 管理系统可独立部署，通过内网API调用 Yuxi 和 Python AI服务。")

    doc.add_heading("17. 项目里程碑与实施计划", level=1)
    add_table(
        doc,
        ["阶段", "周期建议", "交付物"],
        [
            ("阶段1：知识库与文献分析MVP", "4-6周", "Yuxi知识库、文献入库、检索问答、翻译、临床字段抽取样例、综述生成样例。"),
            ("阶段2：AI配方推荐与研发建议原型", "4-6周", "原料图谱、保湿配方推荐List、研发方向识别、假设生成、证据链、实验建议、人工复核流程。"),
            ("阶段3：中试模拟原型", "6-8周", "小试数据录入、放大规则、设备适配、方案生成、反馈回填、误差分析。"),
            ("阶段4：Java管理系统一期", "8-10周", "用户权限、项目任务、文献任务、配方候选、实验方案、原料/设备供应商、采购协同。"),
            ("阶段5：虚实转换试点", "视设备条件", "设备数据采集接口、字段映射、方案回写、审批和真实数据回流。"),
        ],
        [1700, 1500, 6160],
        LIGHT_BLUE,
    )

    doc.add_heading("18. 风险、难点与应对策略", level=1)
    add_table(
        doc,
        ["风险/难点", "影响", "应对策略"],
        [
            ("文献字段不统一", "抽取结果难以直接入库。", "临床文献用模板，科研文献用动态实体关系加人工确认。"),
            ("数值、单位、对象绑定错误", "后续建议和模型训练失真。", "Schema校验、单位归一、证据片段展示、抽样复核。"),
            ("AI幻觉和无依据建议", "影响研发判断。", "强制引用证据链，无证据建议不得进入方案。"),
            ("配方推荐不可解释", "工程师难以信任推荐结果。", "每个原料和比例建议必须给出功效逻辑、图谱关系、历史反馈和风险说明。"),
            ("实验反馈质量不稳定", "AI进化方向可能偏差。", "反馈表单标准化，区分失败原因、工艺问题、原料问题和评价偏差。"),
            ("中试数据量不足", "模型难以稳定优化。", "先用规则和统计模型，逐步积累真实反馈数据。"),
            ("真实设备接口复杂", "虚实转换落地周期长。", "第一版只做接口规范和数据采集，控制指令后置。"),
            ("外部模型数据安全", "企业敏感数据泄露风险。", "脱敏、私有化模型、密钥隔离和调用审计。"),
        ],
        [2200, 3000, 4160],
        LIGHT_GRAY,
    )

    doc.add_heading("19. 验收标准", level=1)
    add_bullets(doc, [
        "文献入库：支持PDF/DOCX文献上传、解析、分块、检索和引用来源展示。",
        "临床抽取：能够按模板抽取关键字段，并保留原文证据和翻译结果。",
        "科研归纳：能够围绕指定方向生成可追溯综述，文献链接可访问或下载。",
        "专利分析：能够输出专利摘要、权利要求拆解、相似点和差异点。",
        "研发建议：每条建议包含证据、假设、实验验证方法和失败条件。",
        "配方推荐：输入“保湿配方”等目标后，系统能够输出多套候选配方List，包含原料组成、比例范围、推荐理由、风险提示、证据链和推荐评分。",
        "工程师筛选：配方工程师能够对AI推荐配方进行保留、剔除、修改和备注，系统保存筛选原因。",
        "实验反馈学习：真实实验结果能够回写到配方和原料关系中，并影响后续同类配方推荐排序。",
        "采购衔接：验证可行配方能够关联原材料供应商、样品采购状态、价格、交期和资质信息。",
        "中试模拟：能够录入小试数据，生成至少一套中试参数建议、风险清单和设备候选。",
        "反馈闭环：能够回填线下实验结果，并生成预测-实测差异分析。",
        "权限审计：关键操作有用户、时间、对象、前后状态和来源记录。",
    ])

    doc.add_heading("20. 后续演进路线", level=1)
    add_para(doc, "第一版完成后，系统可沿四个方向演进：一是扩大知识源和领域抽取模板，形成企业研发知识资产；二是增强原料图谱和配方推荐模型，让实验反馈持续改变推荐权重；三是增强中试模拟算法，引入更多真实反馈、机理模型和参数优化；四是建设 Java 管理系统和工厂接口，使AI建议真正进入企业流程和生产数据闭环。")
    add_callout(doc, "最终愿景", "把系统从“会总结文献的AI工具”推进到“能基于原料图谱推荐配方、接受工程师筛选和实验反馈、持续进化并衔接采购和中试的配方研发操作系统”。")

    doc.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    build()
